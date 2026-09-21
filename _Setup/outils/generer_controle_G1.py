#!/usr/bin/env python3
"""generer_controle_G1.py — LOT 1, contrôle d'ouverture G1 (FON-S7, TD Hélice marine).

Produit les deux DOCX à la charte ENSM à partir des sources Markdown
`Helice/docs/13_CONTROLE-G1_SUJET.md` et `14_CONTROLE-G1_ENSEIGNANT.md` (contenu déjà
rédigé et relu contradictoirement — LOT 0B). Ce script ne récrit pas le contenu : il le
met en forme selon la consigne LOT 1 (format « carré » du prédécesseur).

Base : `~/ENSM-Enseignement/_Setup/gabarits/reference_ENSM_polycopie_CHARTE.docx` (charte
officielle, chargée en lecture seule — jamais modifiée). Assets d'identité chargés depuis
`~/ENSM-Enseignement/_Setup/identite/` (jamais modifiés).

Composition de l'en-tête de première page : la vague en pleine largeur, puis un tableau
sans bordure (logo à gauche, date + nom à droite) juste en dessous — PAS un chevauchement
pixel-exact (recadrage/positionnement flottant en OpenXML n'est pas vérifiable dans cet
environnement : LibreOffice ne charge plus aucun DOCX, cf. INV-22 / porte LOT 1.3). C'est le
même schéma de composition, déjà en production, que `page_de_garde_ENSM.py` (vague, puis
bloc titre, empilés) — pas une invention pour ce script.

Usage :
    python3 generer_controle_G1.py
"""
import os
import re
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ENSM_ENS = os.path.expanduser("~/ENSM-Enseignement")
REFERENCE = os.path.join(ENSM_ENS, "_Setup/gabarits/reference_ENSM_polycopie_CHARTE.docx")
IDENTITE = os.path.join(ENSM_ENS, "_Setup/identite")
VAGUE = os.path.join(IDENTITE, "ENSM_vague_page_de_garde.png")
LOGO = os.path.join(IDENTITE, "ENSM_logo_couleur.png")

ICI = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.join(ICI, "..", "..", "Helice", "docs")

MARINE = RGBColor(0x1A, 0x34, 0x6D)
TEAL = RGBColor(0x1A, 0x99, 0x88)

RUNNING_HEADER = "Contrôle d'ouverture – Environnement de travail – I4-S7-FON – 07/09/2026"


# --------------------------------------------------------------------------- helpers

def add_field(paragraph, instr, bold=False, size=None, color=None):
    """Insère un champ Word (ex. PAGE, NUMPAGES) dans un paragraphe existant."""
    def _mkrun():
        r = paragraph.add_run()
        if bold:
            r.bold = True
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = color
        return r

    r1 = _mkrun()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    r1._r.append(fld_begin)

    r2 = _mkrun()
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = instr
    r2._r.append(instrText)

    r3 = _mkrun()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    r3._r.append(fld_sep)

    r4 = _mkrun()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    r4._r.append(fld_end)


def set_cell_borders_none(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    tcPr.append(borders)


def bottom_rule(paragraph, color="1A9988", size=10, space=6):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), str(space))
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def clear_body(document):
    body = document.element.body
    sectPr = body.find(qn('w:sectPr'))
    for child in list(body):
        if child is not sectPr:
            body.remove(child)



def clear_hf(hf):
    """Vide un en-tête/pied de page hérité du reference-doc (contenu existant :
    logo, "Page X | Y") avant d'y écrire un contenu neuf -- sinon les runs/images
    d'origine restent et se mélangent avec le nouveau texte."""
    for p in list(hf.paragraphs):
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        if p is not hf.paragraphs[0]:
            p._p.getparent().remove(p._p)
    for t in list(hf.tables):
        t._tbl.getparent().remove(t._tbl)
    if not hf.paragraphs:
        hf.add_paragraph()


def build_first_page_header(section, mention_enseignant=False):
    section.different_first_page_header_footer = True
    hdr = section.first_page_header
    hdr.is_linked_to_previous = False
    clear_hf(hdr)
    p_img = hdr.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_after = Pt(0)
    p_img.add_run().add_picture(VAGUE, width=Cm(16.6), height=Cm(1.7))  # bandeau resserre (LOT1 : sacrifice hauteur bandeau avant contenu)

    tbl = hdr.add_table(rows=1, cols=2, width=Cm(16.6))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    left, right = tbl.rows[0].cells
    set_cell_borders_none(left)
    set_cell_borders_none(right)

    left.paragraphs[0].add_run().add_picture(LOGO, width=Cm(4.2))

    p1 = right.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run("07/09/2026")
    r1.bold = True
    r1.font.color.rgb = MARINE
    r1.font.size = Pt(11)

    p2 = right.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run("M. Christophe Vanhorick")
    r2.font.color.rgb = MARINE
    r2.font.size = Pt(11)


def build_running_header(section):
    hdr = section.header
    hdr.is_linked_to_previous = False
    clear_hf(hdr)
    p = hdr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bottom_rule(p)
    r = p.add_run(RUNNING_HEADER)
    r.font.size = Pt(9)
    r.font.color.rgb = MARINE


def build_footer(section):
    ftr = section.footer
    ftr.is_linked_to_previous = False
    clear_hf(ftr)
    p = ftr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE", size=Pt(9), color=MARINE)
    r = p.add_run(" / ")
    r.font.size = Pt(9)
    r.font.color.rgb = MARINE
    add_field(p, "NUMPAGES", size=Pt(9), color=MARINE)

    ftr_fp = section.first_page_footer
    ftr_fp.is_linked_to_previous = False
    clear_hf(ftr_fp)
    p2 = ftr_fp.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p2, "PAGE", size=Pt(9), color=MARINE)
    r2 = p2.add_run(" / ")
    r2.font.size = Pt(9)
    r2.font.color.rgb = MARINE
    add_field(p2, "NUMPAGES", size=Pt(9), color=MARINE)


def add_titre_bloc(document, mention_enseignant):
    p1 = document.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("I4 – S7 – FON – Mécanique des fluides / Hydrodynamique")
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = MARINE

    p2 = document.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("Contrôle d'ouverture — Environnement de travail")
    r2.bold = True
    r2.underline = True
    r2.font.size = Pt(12.5)
    r2.font.color.rgb = MARINE

    if mention_enseignant:
        p3 = document.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(6)
        r3 = p3.add_run("DOCUMENT ENSEIGNANT — NE PAS DIFFUSER")
        r3.bold = True
        r3.font.size = Pt(10)
        r3.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)


def add_modalites(document, nb_pages_sujet="1"):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Modalités de l'épreuve :")
    r.bold = True
    r.font.size = Pt(9.5)

    items = [
        ("Durée", "30 minutes"),
        ("Points", "20 pts (+ 2 pts bonus possibles, question B3)"),
        ("Documents autorisés", "aucun"),
        ("Type d'épreuve", "contrôle écrit individuel"),
        ("Langue", "Français"),
        ("Nombre de pages du sujet", nb_pages_sujet),
    ]
    for label, valeur in items:
        bp = document.add_paragraph()
        bp.paragraph_format.left_indent = Cm(0.5)
        bp.paragraph_format.space_after = Pt(0)
        bp.add_run("•  ").bold = True
        rl = bp.add_run(label)
        rl.bold = True
        rl.underline = True
        rl.font.size = Pt(9.5)
        rv = bp.add_run(f" : {valeur}")
        rv.bold = True
        rv.font.size = Pt(9.5)

    bp = document.add_paragraph()
    bp.paragraph_format.left_indent = Cm(0.5)
    bp.paragraph_format.space_after = Pt(0)
    bp.add_run("•  ").bold = True
    bp.runs[-1].font.size = Pt(9.5)
    rl = bp.add_run("Remarques")
    rl.bold = True
    rl.underline = True
    rl.font.size = Pt(9.5)
    rl2 = bp.add_run(" :")
    rl2.bold = True
    rl2.font.size = Pt(9.5)

    remarques = [
        "Pensez à bien indiquer vos nom(s) et prénom(s) sur TOUTES les copies de réponse ;",
        "Les trois parties sont indépendantes : traitez-les dans l'ordre que vous voulez ;",
        "On note le raisonnement, pas l'orthographe des commandes — une faute de frappe ou "
        "un accent manquant ne coûte rien ;",
        "Ce contrôle porte sur l'environnement de travail installé au semestre dernier, pas sur le "
        "TD qui commence aujourd'hui.",
    ]
    for rem in remarques:
        sp = document.add_paragraph()
        sp.paragraph_format.left_indent = Cm(1.1)
        sp.paragraph_format.space_after = Pt(0)
        sp.add_run("o  ").font.size = Pt(9)
        rr = sp.add_run(rem)
        rr.font.size = Pt(9)

    pc = document.add_paragraph()
    pc.paragraph_format.space_before = Pt(6)
    pc.paragraph_format.space_after = Pt(2)
    rc = pc.add_run("Bon courage !")
    rc.italic = True
    rc.font.size = Pt(9.5)

    hr = document.add_paragraph()
    hr.paragraph_format.space_after = Pt(4)
    bottom_rule(hr, color="1A346D", size=6, space=1)


PART_RE = re.compile(r'^# Partie ([A-Z]) — (.+?) \*\((\d+(?:,\d+)? points?), ~(\d+) min\)\*$')
QUESTION_RE = re.compile(r'^\*\*([A-Z]\d) — (.+?) \*\((\d+(?:,\d+)?) pts?\)\*\.\*\*\s*(.*)$')
SUBQ_RE = re.compile(r'^\*\*([a-c])\)\*\*\s*(.+)$')

ROMANS = {"A": "I", "B": "II", "C": "III"}


def strip_md_bold(text):
    return re.sub(r'\*\*(.+?)\*\*', r'\1', text)


def add_wrapped(paragraph, text, base_size=Pt(9.5)):
    """Ajoute `text` (avec **gras** Markdown) comme runs dans un paragraphe existant."""
    pos = 0
    for m in re.finditer(r'\*\*(.+?)\*\*', text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.font.size = base_size
        rb = paragraph.add_run(m.group(1))
        rb.bold = True
        rb.font.size = base_size
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.font.size = base_size


def parse_and_add_body(document, md_lines):
    """Convertit les lignes du corps (après le bloc Modalités) en paragraphes Word,
    en respectant le format imposé (parties en chiffres romains, questions numérotées,
    sous-questions a/b/c)."""
    i = 0
    while i < len(md_lines):
        line = md_lines[i].rstrip("\n")
        i += 1
        if not line.strip():
            continue

        m = PART_RE.match(line)
        if m:
            lettre, titre, pts, minutes = m.groups()
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f"{ROMANS[lettre]} – {titre}")
            r.bold = True
            r.underline = True
            r.font.size = Pt(11)
            r.font.color.rgb = MARINE
            r2 = p.add_run(f"  ({pts}, ~{minutes} min)")
            r2.font.size = Pt(8.5)
            r2.italic = True
            continue

        m = QUESTION_RE.match(line)
        if m:
            code, titre, pts, reste = m.groups()
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(f"{code} — {titre}")
            r.bold = True
            r.underline = True
            r.font.size = Pt(9.5)
            r2 = p.add_run(f" : ({pts} pt{'s' if pts != '1' else ''})")
            r2.bold = True
            r2.font.size = Pt(9.5)
            if reste.strip():
                p2 = document.add_paragraph()
                p2.paragraph_format.space_after = Pt(2)
                add_wrapped(p2, reste.strip())
            continue

        m = SUBQ_RE.match(line)
        if m:
            lettre, texte = m.groups()
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(f"{lettre}. ")
            r.bold = True
            r.font.size = Pt(9.5)
            add_wrapped(p, texte)
            continue

        if line.startswith("```"):
            # bloc de code (arborescence, message d'erreur) : police monospace
            block = []
            while i < len(md_lines) and not md_lines[i].startswith("```"):
                block.append(md_lines[i].rstrip("\n"))
                i += 1
            i += 1
            for bl in block:
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.6)
                r = p.add_run(bl if bl else " ")
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
            continue

        # paragraphe normal
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_wrapped(p, line.strip())

    return document


def read_source(md_path, marker="# Partie A"):
    with open(md_path, encoding="utf-8") as f:
        text = f.read()
    if text.lstrip().startswith("---"):
        text = text.split("---", 2)[2]
    lines = text.split("\n")
    start = 0
    for idx, l in enumerate(lines):
        if l.startswith(marker):
            start = idx
            break
    return lines[start:]


# ------------------------------------------------------------- rendu générique (doc 14)
# `14_CONTROLE-G1_ENSEIGNANT.md` est un document de travail riche (relecture, corrigé,
# barème, banque orale, problématiques) : "pagination libre" (LOT 1). Pas de maquette
# stricte à reproduire ici, juste une lecture confortable pour l'enseignant -- rendu
# markdown générique (titres, listes, tableaux, citations, gras/italique), pas la mise
# en page « carré » du sujet.

def set_table_borders(table, color="A9D5E2"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def render_table(document, rows):
    header, *body = rows
    tbl = document.add_table(rows=1, cols=len(header))
    set_table_borders(tbl)
    for cell, text in zip(tbl.rows[0].cells, header):
        p = cell.paragraphs[0]
        add_wrapped(p, text.strip(), base_size=Pt(9.5))
        for r in p.runs:
            r.bold = True
    for row in body:
        cells = tbl.add_row().cells
        for cell, text in zip(cells, row):
            add_wrapped(cell.paragraphs[0], text.strip(), base_size=Pt(9.5))


def render_markdown_generic(document, lines):
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        i += 1

        if not stripped or stripped == "---":
            continue

        if stripped.startswith("```"):
            block = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for bl in block:
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.6)
                r = p.add_run(bl if bl.strip() else " ")
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            continue

        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt({1: 15, 2: 13, 3: 11.5, 4: 10.5}[level])
            r.font.color.rgb = MARINE
            if level <= 2:
                bottom_rule(p, color="A9D5E2", size=6, space=2)
            continue

        if stripped.startswith(">"):
            block = [stripped.lstrip(">").strip()]
            while i < n and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            left.set(qn('w:val'), 'single')
            left.set(qn('w:sz'), '10')
            left.set(qn('w:space'), '6')
            left.set(qn('w:color'), '1A9988')
            pBdr.append(left)
            pPr.append(pBdr)
            add_wrapped(p, " ".join(block), base_size=Pt(9.5))
            for r in p.runs:
                r.italic = True
            continue

        if stripped.startswith("|"):
            rows = [stripped]
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            parsed = []
            for ridx, row in enumerate(rows):
                if ridx == 1 and re.match(r'^\|[\s:|-]+\|$', row):
                    continue
                cells = [c.strip() for c in row.strip("|").split("|")]
                parsed.append(cells)
            if parsed:
                render_table(document, parsed)
            continue

        m = re.match(r'^[-*]\s+(.*)$', stripped)
        if m:
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(0)
            p.add_run("•  ")
            add_wrapped(p, m.group(1), base_size=Pt(10))
            continue

        m = re.match(r'^\d+\.\s+(.*)$', stripped)
        if m:
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(0)
            p.add_run(m.group(0).split(".",1)[0] + ".  ")
            add_wrapped(p, m.group(1), base_size=Pt(10))
            continue

        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_wrapped(p, stripped, base_size=Pt(10))


def build(md_path, out_path, mention_enseignant, margin_cm=None, strict=True, nb_pages="1"):
    doc = Document(REFERENCE)
    clear_body(doc)
    section = doc.sections[0]
    if margin_cm:
        section.top_margin = section.bottom_margin = Cm(margin_cm)
        section.left_margin = section.right_margin = Cm(margin_cm)

    build_first_page_header(section, mention_enseignant)
    build_running_header(section)
    build_footer(section)

    add_titre_bloc(doc, mention_enseignant)

    if strict:
        add_modalites(doc, nb_pages_sujet=nb_pages)
        lines = read_source(md_path, marker="# Partie A")
        parse_and_add_body(doc, lines)
    else:
        lines = read_source(md_path, marker="## 0.")
        render_markdown_generic(doc, lines)

    doc.save(out_path)
    print("écrit :", out_path)


if __name__ == "__main__":
    sujet_md = os.path.join(DOCS, "13_CONTROLE-G1_SUJET.md")
    ens_md = os.path.join(DOCS, "14_CONTROLE-G1_ENSEIGNANT.md")

    build(sujet_md, os.path.join(DOCS, "13_CONTROLE-G1_SUJET.docx"),
          mention_enseignant=False, margin_cm=2.2, strict=True, nb_pages="2")

    # Le document enseignant est un document de travail riche (relecture, corrigé,
    # barème, banque orale, problématiques) : "pagination libre" (LOT 1), pas de
    # maquette stricte à reproduire. Même bandeau et bloc titre que le sujet
    # (+ mention), puis rendu markdown générique pour le reste (titres, listes,
    # tableaux, citations) -- lisible, pas maquetté à la charte "carrée".
    build(ens_md, os.path.join(DOCS, "14_CONTROLE-G1_ENSEIGNANT.docx"),
          mention_enseignant=True, margin_cm=None, strict=False)
